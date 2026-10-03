"""Offline product production only; never called by the Cloudflare runtime."""
from pathlib import Path
from datetime import datetime
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

ROOT = Path(__file__).resolve().parents[1]
DATE = "2026-10-03"
VERSION = "1.0"
NOTICE = "Alat bantu operasional; bukan software, nasihat legal/medis/pajak, atau jaminan pendapatan, pertumbuhan dan retensi."
BASE = ["00-README.pdf", "01-SOP-Dasar.docx", "02-Menu-dan-Daftar-Harga.docx", "03-Checklist-Buka-Tutup.docx", "04-Follow-Up-Pelanggan-Berizin.docx"]
EXTRA = ["05-Rekap-Kunjungan.xlsx", "06-Tracker-Repeat-Visit.xlsx", "07-Review-Bulanan.docx", "08-Panduan-Follow-Up-dan-Review.docx"]
SYSTEM = ["09-Peta-Operating-System.docx", "10-Peta-Customer-Journey.docx", "11-Peta-Retention-dan-Proses.docx", "12-Owner-Review-Framework.docx", "13-Prioritas-Implementasi.docx"]
PACKAGES = [("Starter", "starter", "RBS-STARTER-001", 39000, BASE), ("Growth", "growth", "RBS-GROWTH-001", 79000, BASE + EXTRA), ("System", "system", "RBS-SYSTEM-001", 149000, BASE + EXTRA + SYSTEM)]

# Distinct, usable instructions and editable operational tables for each component.
CONTENT = {
"01-SOP-Dasar.docx": ("SOP dasar layanan barber", [
("Mulai dengan satu alur", "Isi [NAMA USAHA], [HARI/JAM BUKA] dan [PIC PERAN]. Sesuaikan langkah dengan alat serta aturan kerja setempat. Uji pada satu giliran, catat hambatan, lalu revisi. Jangan memakai contoh ini sebagai standar medis atau sertifikasi keselamatan."),
("Urutan layanan", "Sambut pelanggan → pahami layanan dan preferensi → jelaskan menu/harga usaha → konfirmasi pilihan → lakukan layanan sesuai kompetensi operator → konfirmasi hasil → pencatatan sederhana → rapikan area. Jangan menyamakan harga layanan barber pada template dengan harga paket Ranel."),
("Ketika terjadi masalah", "Hentikan langkah yang tidak aman. Dengarkan keluhan, ulangi pemahaman tanpa menyalahkan, lalu arahkan ke PIC. Jangan menjanjikan remedy di luar kewenangan atau hasil yang tidak bisa dipastikan.")],
[(["Tahap", "Langkah yang disepakati", "PIC", "Bukti selesai"], [["Buka", "Periksa area, alat dan menu", "[PIC]", "Checklist"], ["Sambut", "Tanyakan kebutuhan dan pilihan", "[PIC]", "Pilihan dikonfirmasi"], ["Layanan", "Jalankan layanan yang disepakati", "[PIC]", "Hasil dikonfirmasi"], ["Tutup", "Rapikan dan tinjau hambatan", "[PIC]", "Catatan review"]]), (["Masalah operasional", "Tindakan/PIC", "Tanggal review"], [["[Isi tanpa identitas pelanggan]", "[Tindakan]", "[Tanggal]"]])]),
"02-Menu-dan-Daftar-Harga.docx": ("Menu dan daftar harga usaha", [
("Cara mengisi", "Gunakan layanan dan harga yang benar-benar supplied oleh pemilik usaha. Kolom sengaja kosong; tidak ada harga barber resmi yang ditentukan Ranel. Isi harga sebelum dipakai kepada pelanggan, dan perbarui tanggal versi."),
("Sebelum layanan", "Pastikan pelanggan memahami layanan, harga dan tambahan yang dipilih. Jangan menambahkan layanan berbayar tanpa persetujuan. Tempatkan menu pada lokasi yang dapat dibaca pelanggan.")],
[(["Nama layanan", "Deskripsi/batas", "Harga IDR dari pemilik", "Durasi perkiraan"], [["[LAYANAN 1]", "[Cakupan]", "[HARGA PEMILIK]", "[Estimasi]"], ["[LAYANAN 2]", "[Cakupan]", "[HARGA PEMILIK]", "[Estimasi]"]]), (["Hari/jam", "Kontak usaha", "Berlaku mulai", "Disetujui oleh peran"], [["[HARI/JAM]", "[KANAL KONTAK]", "[Tanggal]", "[PIC]"]])]),
"03-Checklist-Buka-Tutup.docx": ("Checklist buka dan tutup", [
("Penggunaan harian", "Duplikasi tabel untuk setiap hari/giliran. Tulis tanggal dan PIC peran. Isi Ya/Tidak/Tidak relevan; tidak perlu menyimpan nama pelanggan. Temuan tidak selesai dibawa ke review, bukan ditandai selesai."),
("Review kecil", "Pilih satu item yang sering terlewat. Tentukan perbaikan, pemilik tindakan dan kapan akan dicek; jangan menambah checklist tanpa kegunaan.")],
[(["Buka: pemeriksaan", "Status", "Temuan/tindakan", "PIC"], [[x, "[Ya/Tidak/N/A]", "[Temuan]", "[PIC]"] for x in ["Area siap dan tidak ada hambatan", "Alat diperiksa sesuai SOP usaha", "Menu/harga yang benar terlihat", "Jam/kanal kontak benar", "Tugas giliran dibagi"]]), (["Tutup: pemeriksaan", "Status", "Temuan/tindakan", "PIC"], [[x, "[Ya/Tidak/N/A]", "[Temuan]", "[PIC]"] for x in ["Area dan alat dirapikan sesuai aturan usaha", "Catatan layanan direkap", "Perlengkapan yang perlu diganti dicatat", "Masalah dibawa ke PIC", "Persiapan giliran berikutnya jelas"]])]),
"04-Follow-Up-Pelanggan-Berizin.docx": ("Follow-up pelanggan berbasis izin", [
("Izin terlebih dahulu", "Tanyakan apakah pelanggan ingin dihubungi, tujuan dan kanalnya. Penolakan tidak mengurangi layanan. Pisahkan reminder dari promo. Catat izin secara minimal di tempat pribadi milik usaha; jangan mengirim daftar customer atau nomor telepon kepada Ranel."),
("Contoh teks, bukan otomasi", "Permintaan izin: Apakah Anda bersedia menerima [TUJUAN PESAN] dari [NAMA USAHA] melalui [KANAL]? Anda dapat berhenti kapan saja. Pesan berizin: Halo, ini [NAMA USAHA]. Sesuai izin Anda, kami mengirim [TUJUAN]. Bila tidak ingin menerima lagi, beri tahu kami. Sesuaikan wording; jangan kirim bila izin tidak ada."),
("Batas dan penghentian", "Sebelum mengirim cek tujuan, izin dan apakah ada permintaan berhenti. Jangan blast, mengambil kontak publik tanpa izin, atau mengarang riwayat kunjungan. Tindak lanjut dilakukan manual oleh usaha, bukan oleh Ranel.")],
[(["Kode lokal anonim", "Tujuan/kanal disetujui", "Tanggal izin", "Status berhenti", "PIC"], [["[KODE LOKAL]", "[Tujuan/kanal]", "[Tanggal]", "[Ya/Tidak]", "[PIC]"]]), (["Langkah", "Cek sebelum kirim"], [["Izin", "Masih berlaku untuk tujuan ini?"], ["Isi", "Tidak sensitif dan relevan?"], ["Frekuensi", "Sesuai kesepakatan customer?"], ["Stop", "Permintaan berhenti sudah diterapkan?"]])]),
"07-Review-Bulanan.docx": ("Review bulanan sederhana", [
("Gunakan data yang benar", "Rekap angka dari workbook usaha sendiri. Sampel bukan revenue nyata. Total nilai layanan bukan laba atau ledger pembayaran. Pisahkan data kosong dari nol; tulis keterbatasan pencatatan."),
("Pilih satu eksperimen", "Tinjau penggunaan checklist, jenis layanan, kunjungan dan repeat-visit. Catat satu hambatan yang terbukti, satu tindakan realistis dan tanggal review berikut. Jangan menyimpulkan pertumbuhan karena perubahan kecil tanpa pembanding.")],
[(["Periode", "Jumlah kunjungan tercatat", "Nilai layanan", "Keterbatasan data"], [["[Bulan]", "[Dari workbook]", "[Dari workbook]", "[Catatan]"]]), (["Apa diamati", "Bukti", "Satu tindakan", "PIC", "Review"], [["[Masalah]", "[Bukti agregat]", "[Tindakan]", "[PIC]", "[Tanggal]"]])]),
"08-Panduan-Follow-Up-dan-Review.docx": ("Panduan follow-up dan review", [
("Rutinitas Growth", "Isi Input pada Rekap Kunjungan setelah layanan. Gunakan kode customer lokal bila memang dibutuhkan dan diizinkan. Pada Tracker, kode yang sama menunjukkan kunjungan berulang dalam data yang dicatat, bukan pelanggan unik secara otomatis."),
("Baca workbook", "Ringkasan menghitung baris/nilai yang tercatat, bukan semua aktivitas usaha. Persentase repeat adalah kode dengan minimal dua kunjungan dibagi kode yang ditinjau; bukan cohort retention dan tidak mengukur izin marketing. Demo berisi data sintetis; jangan salin sebagai transaksi nyata."),
("Sebelum follow-up", "Cek izin di catatan privat usaha; workbook hanya membantu review. Pilih tujuan, lakukan pesan manual, catat outcome agregat. Jangan mengimpor database customer ke Ranel atau menganggap checklist sebagai otomasi."),
("Bulanan", "Isi Review Bulanan. Pilih tindakan berdasarkan bukti dan kapasitas; tinjau hasil dan keterbatasan. Tidak ada jaminan tambahan kunjungan atau pendapatan.")],
[(["Rutinitas", "Peran", "Selesai/isu"], [["Catat layanan", "[PIC]", "[Catatan]"], ["Cek izin sebelum pesan", "[PIC]", "[Catatan]"], ["Review agregat", "[Owner]", "[Catatan]"]])]),
"09-Peta-Operating-System.docx": ("Peta operating system usaha", [
("Peta, bukan aplikasi", "System mencakup Growth dan dokumen pemetaan kerja. Tuliskan proses yang benar-benar terjadi. Tidak ada software, hosting, integrasi atau engineering commitment dalam paket."),
("Cara memakai", "Mulai dari satu alur utama. Tentukan input, langkah, output, peran dan bukti selesai. Jika sebuah proses belum ada, tulis Belum ada; jangan mengisi seolah sudah berjalan.")],
[(["Alur", "Input", "Langkah/output", "Peran", "Bukti"], [[x, "[Input]", "[Langkah/output]", "[PIC]", "[Bukti]"] for x in ["Buka", "Layanan", "Catatan", "Follow-up berizin", "Tutup", "Review"]])]),
"10-Peta-Customer-Journey.docx": ("Peta perjalanan pelanggan", [
("Pemetaan tanpa data pribadi", "Catat tahapan dan hambatan secara umum, bukan nama/telepon customer. Cek apa yang customer perlu pahami pada tiap tahap. Data private tetap di usaha, tidak dikirim ke Ranel."),
("Tinjau satu titik", "Pilih titik yang menimbulkan kebingungan. Uji perubahan kecil dan catat bukti; jangan mengklaim journey baru meningkatkan conversion tanpa data.")],
[(["Tahap", "Customer perlu tahu", "Titik kontak", "Hambatan", "Perbaikan/PIC"], [[x, "[Informasi]", "[Kanal]", "[Hambatan]", "[Aksi]"] for x in ["Mengenal usaha", "Memilih layanan", "Datang", "Menerima layanan", "Konfirmasi hasil", "Kembali bila ingin"]])]),
"11-Peta-Retention-dan-Proses.docx": ("Peta retensi dan proses", [
("Retensi bukan jaminan", "Peta ini menghubungkan pengalaman layanan, izin follow-up dan review kunjungan. Bedakan repeat-visit yang tercatat dari retensi yang disebabkan tindakan tertentu; jangan menganggap hubungan sebagai sebab."),
("Batas izin", "Customer yang menolak atau mencabut izin tidak masuk alur pesan. Informasi kontak dan consent disimpan privat oleh usaha. Tidak ada automated CRM atau pesan Ranel.")],
[(["Proses", "Pemicu", "Syarat izin", "PIC", "Bukti/outcome agregat"], [["Pengalaman layanan", "Layanan selesai", "Tidak meminta kontak tambahan", "[PIC]", "[Catatan]"], ["Follow-up manual", "Tujuan relevan", "Izin masih berlaku", "[PIC]", "[Catatan]"], ["Stop", "Customer meminta berhenti", "Hentikan pesan", "[PIC]", "[Status]"], ["Review repeat", "Periode review", "Data minimal berwenang", "[Owner]", "[Agregat]"]])]),
"12-Owner-Review-Framework.docx": ("Kerangka review pemilik", [
("Agenda pendek", "Tinjau apa yang dijanjikan, apa yang benar-benar berjalan dan hambatan terbesar. Bedakan keputusan owner dari tugas harian. Gunakan data agregat dan jangan mengunggah catatan customer ke Ranel."),
("Keputusan yang dapat ditelusuri", "Untuk setiap tindakan catat alasan, PIC, effort, tanggal review dan hasil sebenarnya. Jika data belum cukup, catat uncertainty dan tunda keputusan besar.")],
[(["Area", "Fakta/bukti", "Gap", "Keputusan", "PIC/review"], [[x, "[Bukti]", "[Gap]", "[Aksi/tunda]", "[PIC/tanggal]"] for x in ["Rutinitas", "Layanan/menu", "Catatan", "Follow-up izin", "Beban kerja"]])]),
"13-Prioritas-Implementasi.docx": ("Prioritas perbaikan praktis", [
("Urutkan sebelum membangun", "Daftar masalah yang memiliki bukti. Bandingkan urgensi, jumlah proses terdampak, effort, risiko dan kapasitas. Mulai dengan perubahan manual yang dapat diuji; dokumen ini bukan backlog software yang otomatis disetujui."),
("Uji dan review", "Pilih satu tindakan, tetapkan bukti selesai dan tanggal tinjau. Jangan mengaktifkan otomasi/pembayaran/integrasi hanya karena ada ide. Perubahan besar memerlukan keputusan pemilik tersendiri.")],
[(["Masalah/bukti", "Urgensi rendah/sedang/tinggi", "Effort", "Risiko", "Aksi/PIC", "Review"], [["[Masalah]", "[Urgensi]", "[Effort]", "[Risiko]", "[Aksi/PIC]", "[Tanggal]"] for _ in range(3)])]),
}


def metadata(name, pid, sku):
    return f"Ranel Barber {name} | {pid} | {sku} | Versi {VERSION} | Disiapkan {DATE}"


def make_doc(path, name, pid, sku):
    title, sections, tables = CONTENT[path.name]
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.65)
    style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(7)
    doc.core_properties.author = 'Ranel'; doc.core_properties.title = title
    doc.core_properties.subject = metadata(name, pid, sku)
    doc.add_paragraph(metadata(name, pid, sku), 'Subtitle')
    doc.add_heading(title, 0)
    doc.add_paragraph('[NAMA USAHA] • [HARI/JAM BUKA] • [KANAL KONTAK] • [LOGO OPSIONAL]')
    doc.add_paragraph('Template kosong untuk disesuaikan. Semua placeholder [TEKS] sengaja belum diisi; tidak mewakili customer nyata.')
    for heading, text in sections:
        doc.add_heading(heading, 1); doc.add_paragraph(text)
    for headers, rows in tables:
        table = doc.add_table(rows=1, cols=len(headers)); table.style = 'Light Shading Accent 1'
        for c, h in zip(table.rows[0].cells, headers): c.text = h
        for row in rows:
            for c, value in zip(table.add_row().cells, row): c.text = value
        doc.add_paragraph('')
    doc.add_paragraph('Gunakan untuk usaha sendiri. Jangan menjual ulang, mensublisensikan, membagikan publik, atau mengemas template Ranel sebagai produk sendiri. Informasi usaha yang Anda isi tetap milik Anda.')
    doc.add_paragraph(NOTICE)
    sec.footer.paragraphs[0].text = f'Ranel • v{VERSION} • {DATE} • Review sebelum digunakan'
    doc.save(path)


def format_sheet(ws):
    ws.freeze_panes = 'A7'
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row == 6:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='536145')
    from openpyxl.utils import get_column_letter
    for index in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(index)].width = 23
    ws.sheet_view.showGridLines = True
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:6'


def init_sheet(wb, title, meta, instructions):
    ws = wb.create_sheet(title)
    ws['A1'] = meta; ws['A2'] = 'DATA CONTOH SINTETIS SAJA' if 'Contoh' in title else 'TEMPLATE KOSONG — data tetap di usaha Anda'
    ws['A3'] = instructions; ws['A4'] = NOTICE
    for r in range(1,5): ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    ws.row_dimensions[3].height = 48; ws.row_dimensions[4].height = 34
    return ws


def make_book(path, name, pid, sku):
    wb = Workbook(); wb.remove(wb.active)
    wb.properties.creator = 'Ranel'; wb.properties.title = path.stem
    meta = metadata(name, pid, sku)
    guide = init_sheet(wb,'Panduan',meta,'Isi template Input/Kunjungan, bukan Contoh. Demo tidak boleh dianggap aktivitas nyata. PDF dan Excel/LibreOffice kompatibel diperlukan; tidak ada macros, koneksi eksternal atau upload otomatis.')
    guide.append([''])
    for r, text in enumerate(['[NAMA USAHA] | [HARI/JAM] | [KANAL KONTAK]','Input hanya data minimal. Pakai kode lokal anonim; jangan masukkan telepon/password/database pelanggan.','Kolom formula jangan ditimpa. Blank bukan nol; catatan layanan bukan financial ledger atau profit.','Formula dibatasi 200 baris; untuk memperluas, salin formula dan perluas range ringkasan lalu uji.','Izin follow-up diperiksa terpisah di catatan privat usaha. Kode kunjungan bukan izin promo.','Lisensi penggunaan usaha sendiri; tidak untuk resale/redistribusi. Dukungan/revisi mengikuti kesepakatan.'],6):
        guide.cell(r,1,text); guide.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
        guide.row_dimensions[r].height=40
    if path.name.startswith('05'):
        for title, demo in [('Input',False),('Contoh',True)]:
            ws=init_sheet(wb,title,meta,'Masukkan tanggal (date), layanan, kategori kunjungan, jumlah dan harga layanan dari usaha. Total = jumlah × harga, bukan laba. Baris demo bersifat sintetis.')
            for col,h in enumerate(['Tanggal','Layanan','Kunjungan baru/ulang','Jumlah layanan','Harga layanan IDR','Nilai layanan IDR'],1):ws.cell(6,col,h)
            if demo:
                for r,vals in enumerate([[datetime(2026,10,1),'Layanan Contoh A','Baru',1,25000],[datetime(2026,10,2),'Layanan Contoh B','Ulang',1,30000],[datetime(2026,10,3),'Layanan Contoh A','Ulang',2,25000]],7):
                    for c,v in enumerate(vals,1):ws.cell(r,c,v)
            for r in range(7,207):
                ws.cell(r,6,f'=IF(OR(A{r}="",D{r}="",E{r}=""),"",D{r}*E{r})')
                ws.cell(r,1).number_format='yyyy-mm-dd'
                for c in [5,6]:ws.cell(r,c).number_format='#,##0'
            ws.auto_filter.ref='A6:F206'
            dv=DataValidation(type='list',formula1='"Baru,Ulang"');ws.add_data_validation(dv);dv.add('C7:C206')
        summary=init_sheet(wb,'Ringkasan',meta,'Ringkasan hanya Input, bukan Contoh; jumlah baris = kunjungan yang dicatat. Tidak mewakili seluruh usaha bila Input belum lengkap.')
        for r,(label,formula) in enumerate([('Kunjungan tercatat','=COUNT(Input!A7:A206)'),('Jumlah layanan','=SUM(Input!D7:D206)'),('Nilai layanan IDR','=SUM(Input!F7:F206)'),('Kunjungan ulang tercatat','=COUNTIF(Input!C7:C206,"Ulang")')],6):summary.cell(r,1,label);summary.cell(r,2,formula)
    else:
        for title,demo in [('Kunjungan',False),('ContohKunjungan',True)]:
            ws=init_sheet(wb,title,meta,'Masukkan tanggal dan kode lokal anonim; izin Ya/Tidak adalah catatan operasional, bukan izin otomatis untuk pesan. Jangan memasukkan nama/nomor customer.')
            for c,h in enumerate(['Tanggal','Kode lokal','Izin follow-up Ya/Tidak','Catatan non-sensitif','PIC peran','Status manual'],1):ws.cell(6,c,h)
            if demo:
                for r,vals in enumerate([[datetime(2026,10,1),'DEMO-A','Ya','Contoh sintetis','Operator','Belum ditinjau'],[datetime(2026,10,2),'DEMO-A','Ya','Contoh sintetis','Operator','Belum ditinjau'],[datetime(2026,10,3),'DEMO-B','Tidak','Contoh sintetis','Operator','Tidak dikirimi pesan']],7):
                    for c,v in enumerate(vals,1):ws.cell(r,c,v)
            for r in range(7,207):ws.cell(r,1).number_format='yyyy-mm-dd'
            ws.auto_filter.ref='A6:F206'
            dv=DataValidation(type='list',formula1='"Ya,Tidak"');ws.add_data_validation(dv);dv.add('C7:C206')
        for title,source,demo in [('Tracker','Kunjungan',False),('ContohTracker','ContohKunjungan',True)]:
            ws=init_sheet(wb,title,meta,'Masukkan satu kode unik per baris. Jumlah = COUNTIF kode pada Kunjungan. Repeat berarti ≥2 kunjungan tercatat, bukan jaminan retensi/cohort analytics. Izin dikonfirmasi manual sebelum follow-up.')
            for c,h in enumerate(['Kode lokal unik','Kunjungan tercatat','Repeat Ya/Tidak','Izin terkini (cek manual)','Aksi manual','Review'],1):ws.cell(6,c,h)
            if demo:ws['A7']='DEMO-A';ws['A8']='DEMO-B'
            for r in range(7,207):
                ws.cell(r,2,f'=IF(A{r}="","",COUNTIF({source}!B$7:B$206,A{r}))')
                ws.cell(r,3,f'=IF(A{r}="","",IF(B{r}>=2,"Ya","Tidak"))')
            ws.auto_filter.ref='A6:F206'
        summary=init_sheet(wb,'Ringkasan',meta,'Hanya kode unik yang dimasukkan pada Tracker. Persentase repeat bukan cohort retention dan bukan izin marketing. Tidak ada kontak customer atau pengiriman pesan otomatis.')
        for r,(label,formula) in enumerate([('Kode ditinjau','=COUNTA(Tracker!A7:A206)'),('Kode repeat','=COUNTIF(Tracker!C7:C206,"Ya")'),('Proporsi repeat','=IFERROR(B7/B6,0)')],6):summary.cell(r,1,label);summary.cell(r,2,formula)
        summary['B8'].number_format='0.0%'
    for ws in wb:format_sheet(ws)
    wb.save(path)


def make_pdf(path,name,pid,sku,files,price):
    styles=getSampleStyleSheet(); styles['BodyText'].fontSize=10;styles['BodyText'].leading=15
    story=[]
    def para(t,style='BodyText'):story.extend([Paragraph(t,styles[style]),Spacer(1,10)])
    para(f'Ranel Barber {name}','Title');para(metadata(name,pid,sku),'BodyText')
    para('Paket sumber v1.0 — ASSET_READY_FOR_REVIEW. Bukan bukti pembayaran/fulfillment atau produk Available.','Heading2')
    para(f'Harga dasar approved: Rp{price:,}'.replace(',','.')+' • satu kali • IDR. Checkout belum aktif; payable/tax final mengikuti ketentuan yang berlaku sebelum pembayaran live.')
    para('Mulai menggunakan kit','Heading2')
    for t in ['1. Cocokkan file dengan MANIFEST.md dan baca batas paket.','2. Salin file untuk usaha Anda. Ganti placeholder [NAMA USAHA], [LAYANAN], [HARGA PEMILIK], [HARI/JAM] dan [KANAL KONTAK]; logo opsional.','3. Buka DOCX dengan aplikasi kompatibel DOCX; PDF dengan PDF reader; jika paket menyertakan XLSX gunakan Excel/LibreOffice yang mendukung formula. Tabel dan workbook dapat diedit, tanpa macro/akun/hosted app.','4. Coba satu checklist/rutinitas. Isi data minimal; jangan mengirim password, payment credentials atau database customer kepada Ranel.','5. Sheet Contoh memakai data sintetis, bukan transaksi nyata. Sheet Input/Kunjungan/Tracker kosong. Formula/range sampai 200 baris dijelaskan pada Panduan; jangan menimpa formula.','6. Periksa penerimaan/file-access/manifest match. Jika hilang, rusak atau tidak sesuai scope, hubungi Ranel melalui kontak privat yang disepakati. Waktu delivery/support/revisi harus disepakati; tidak ada SLA baru dari file ini.']:
        para(t)
    para('Manifest komponen','Heading2')
    for f in files:para(f)
    para('Batas, lisensi dan penggunaan','Heading2')
    para('Standard reusable templates + tailoring terbatas. System mencakup seluruh Growth sebagai paket dokumen peta kerja, bukan software. Tidak termasuk custom POS/CRM/booking, hosting, integrasi, unlimited consulting/revisions, atau pesan otomatis.')
    para('Gunakan materi untuk usaha pembeli sendiri. Tidak untuk resale, sublicensing, redistribusi publik atau repackaging. Informasi usaha yang Anda isi tetap milik Anda; reusable template/IP tetap milik Ranel kecuali kesepakatan tertulis lain.')
    para(NOTICE)
    para('Informasi kebijakan: https://ranel.pages.dev/legal • Support privat: https://ranel.pages.dev/contact. Kebijakan/info bukan payment activation. Tidak memerlukan koneksi internet untuk mengedit sumber offline.')
    SimpleDocTemplate(str(path),pagesize=(595.28,841.89),rightMargin=44,leftMargin=44,topMargin=40,bottomMargin=40).build(story)


def main():
    registry=[]
    for name,slug,sku,price,files in PACKAGES:
        pid='ranel.barber.'+slug
        folder=ROOT/'products'/f'Ranel-Barber-{name}-v1.0';folder.mkdir(parents=True,exist_ok=True)
        expected=files+['MANIFEST.md']
        for f in files:
            p=folder/f
            if p.suffix=='.docx':make_doc(p,name,pid,sku)
            elif p.suffix=='.xlsx':make_book(p,name,pid,sku)
            else:make_pdf(p,name,pid,sku,expected,price)
        manifest=f'# Ranel Barber {name} — manifest\n\nProduct ID: `{pid}`\nSKU: `{sku}`\nVersion: {VERSION}\nPreparation date: {DATE}\nAsset state: ASSET_READY_FOR_REVIEW\nSell state: HOLD_PENDING_TERMS\nBase price: IDR {price}, one-time; no checkout activation.\n\n## File canonical\n'+''.join('- `'+f+'`\n' for f in expected)+'\n## Petunjuk dan batas penggunaan\nMulai dari 00-README.pdf. Placeholder sengaja kosong; baris Contoh/demo workbook sintetis, bukan aktivitas nyata. Data usaha tetap di usaha pembeli; jangan mengirim database pelanggan atau credentials kepada Ranel.\n\nGunakan untuk usaha pembeli sendiri; bukan untuk resale, sublicensing, redistribusi publik atau repackaging. System adalah paket dokumen, bukan software.\n\n'+NOTICE+'\n\nDukungan, delivery dan revisi mengikuti kesepakatan yang disetujui; tidak ada SLA publik baru. Produk belum berstatus Available hanya karena file ini ada.\n'
        (folder/'MANIFEST.md').write_text(manifest)
        registry.append({'name':'Ranel Barber '+name,'id':pid,'sku':sku,'version':VERSION,'date':DATE,'price':price,'currency':'IDR','sellState':'HOLD_PENDING_TERMS','assetState':'ASSET_READY_FOR_REVIEW','path':str(folder.relative_to(ROOT)),'files':expected})
    (ROOT/'products/registry.json').write_text(json.dumps(registry,ensure_ascii=False,indent=2)+'\n')
    print('Produced 3 versioned bundles:',sum(len(p[4])+1 for p in PACKAGES),'files. QC required before acceptance.')

if __name__=='__main__':main()
