"""QC real product files with LibreOffice, OOXML parsers and PDF extraction."""
from pathlib import Path
import json, subprocess, shutil, zipfile, hashlib, re
import pymupdf as fitz
from docx import Document
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
QA=ROOT/'qa-artifacts/product-qc'
QA.mkdir(parents=True,exist_ok=True)
registry=json.loads((ROOT/'products/registry.json').read_text())
results=[]

def office(files,out,kind):
    out.mkdir(parents=True,exist_ok=True)
    # A failed conversion must not pass because an older QC file exists.
    for f in files:
        (out/(f.stem+('.pdf' if kind=='pdf' else '.xlsx'))).unlink(missing_ok=True)
    profile=QA/('profile-'+out.name)
    cmd=['libreoffice','-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to',kind,'--outdir',str(out)]+[str(p) for p in files]
    proc=subprocess.run(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,timeout=180)
    if proc.returncode:raise AssertionError('Office conversion failed')
    for f in files:assert (out/(f.stem+('.pdf' if kind=='pdf' else '.xlsx'))).exists(), f.name

for product in registry:
    folder=ROOT/product['path'];out=QA/folder.name
    assert set(p.name for p in folder.iterdir())==set(product['files'])
    manifest=(folder/'MANIFEST.md').read_text()
    listed=re.findall(r'^- `([^`]+)`$',manifest,re.M)
    assert listed==product['files']
    docxs=list(folder.glob('*.docx'));books=list(folder.glob('*.xlsx'))
    if docxs:office(docxs,out/'rendered-docx','pdf')
    if books:
        office(books,out/'recalculated','xlsx')
        for f in books:shutil.copy2(out/'recalculated'/f.name,f)
    for filename in product['files']:
        path=folder/filename;kind=path.suffix[1:]
        if kind in ['docx','xlsx']:
            with zipfile.ZipFile(path) as z:
                assert z.testzip() is None
                assert not any('vbaProject' in n or 'externalLink' in n for n in z.namelist())
        text=''
        if kind=='docx':
            d=Document(path);assert d.tables
            text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
            rendered=fitz.open(out/'rendered-docx'/path.with_suffix('.pdf').name)
            assert len(rendered)>0
            assert sum(len(page.get_text()) for page in rendered)>200
            for page in rendered:
                for x0,y0,x1,y1,*_ in page.get_text('blocks'):
                    assert x0>=-1 and y0>=-1 and x1<=page.rect.width+1 and y1<=page.rect.height+1
        elif kind=='pdf':
            pdf=fitz.open(path);assert len(pdf)>0
            text='\n'.join(p.get_text() for p in pdf)
            assert 'MANIFEST.md' in text and 'bukan' in text
            assert len(text)>1500
            for p in pdf:
                for x0,y0,x1,y1,*_ in p.get_text('blocks'):
                    assert x0>=-1 and y0>=-1 and x1<=p.rect.width+1 and y1<=p.rect.height+1
            pdf[0].get_pixmap(matrix=fitz.Matrix(1,1)).save(str(out/'readme-review.png'))
        elif kind=='xlsx':
            formula=load_workbook(path,data_only=False);cached=load_workbook(path,data_only=True)
            text='\n'.join(str(c.value) for ws in formula for row in ws for c in row if c.value is not None)
            assert any(c.data_type=='f' for ws in formula for row in ws for c in row)
            if filename.startswith('05'):
                assert [cached['Contoh'][f'F{r}'].value for r in [7,8,9]]==[25000,30000,50000]
                assert cached['Ringkasan']['B6'].value==0 and cached['Ringkasan']['B8'].value==0
                assert all(cached['Input'][f'A{r}'].value is None for r in range(7,207))
            else:
                assert cached['ContohTracker']['B7'].value==2 and cached['ContohTracker']['C7'].value=='Ya'
                assert cached['ContohTracker']['B8'].value==1 and cached['ContohTracker']['C8'].value=='Tidak'
                assert cached['Ringkasan']['B6'].value==0 and cached['Ringkasan']['B8'].value==0
                assert all(cached['Kunjungan'][f'A{r}'].value is None for r in range(7,207))
                assert all(cached['Tracker'][f'A{r}'].value is None for r in range(7,207))
            assert not formula._external_links
            # Independently mutate QC copies and recalculate actual office formulas.
            testcopy=out/'formula-test-input'/filename;testcopy.parent.mkdir(parents=True,exist_ok=True)
            if filename.startswith('05'):
                ws=formula['Input'];ws['A7']='2026-10-03';ws['D7']=3;ws['E7']=10000;ws['C7']='Ulang'
            else:
                ws=formula['Kunjungan'];ws['B7']='QC-A';ws['B8']='QC-A';formula['Tracker']['A7']='QC-A'
            formula.save(testcopy);calcdir=out/('formula-test-'+path.stem)
            office([testcopy],calcdir,'xlsx');calc=load_workbook(calcdir/filename,data_only=True)
            if filename.startswith('05'):assert calc['Input']['F7'].value==30000 and calc['Ringkasan']['B8'].value==30000
            else:assert calc['Tracker']['B7'].value==2 and calc['Tracker']['C7'].value=='Ya' and calc['Ringkasan']['B8'].value==1
        else:text=path.read_text()
        normalized=re.sub(r'\s+',' ',text)
        assert product['id'] in normalized and product['sku'] in normalized and '1.0' in normalized and '2026-10-03' in normalized
        assert not re.search(r'(?i)(api[_ -]?key\s*[:=]|NIK\s*[:=]\s*\d|NPWP\s*[:=]\s*\d|wa\.me/\d)',text)
        results.append({'path':str(path.relative_to(ROOT)),'type':kind,'version':product['version'],'qc':'PASS','sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'limitations':'Review template and approved tailoring before delivery; no real customer data or live fulfillment.'})
report={'date':'2026-10-03','state':'ASSET_READY_FOR_REVIEW','files':results,'missing':[], 'checks':['Exact manifests','OOXML opens and editable tables','LibreOffice DOCX to PDF','PDF readability/text bounds','LibreOffice formula recalculation: blank, synthetic, modified QC inputs','No macros/external workbook links','Metadata/version','Targeted privacy checks'],'limitations':['QC templates are synthetic, not a buyer-tailored delivery.','Not legal/tax or provider verification; products remain HOLD_PENDING_TERMS.','Native Excel was not available; formulas were evaluated by LibreOffice.']}
(ROOT/'docs/implementation/phase-02-commerce/product-qc.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print('Product QC PASS:',len(results),'files, 3 bundles. No missing items.')
