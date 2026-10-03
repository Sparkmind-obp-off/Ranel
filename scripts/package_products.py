"""Package already-QC'd private products; crop actual previews without publishing paid sources."""
from pathlib import Path
import json, zipfile, hashlib, shutil, re
import pymupdf as fitz
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
ROOT=Path(__file__).resolve().parents[1]
registry=json.loads((ROOT/'private-products/registry.json').read_text())
qc=json.loads((ROOT/'qa-artifacts/product-maturation-qc.json').read_text())
assert qc['state']=='ASSET_QC_PASS_PENDING_ZIP_AND_VISUAL_REVIEW'
assert len(qc['files'])==31 and not qc['missing']
assert all(f['qc']=='PASS' for f in qc['files'])
for f in qc['files']:assert hashlib.sha256((ROOT/f['path']).read_bytes()).hexdigest()==f['sha256']
release=ROOT/'private-products/releases';release.mkdir(parents=True,exist_ok=True)
preview=ROOT/'public/static/previews';preview.mkdir(parents=True,exist_ok=True)
ledger=[]
font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',20)

def crop(path,page_index,rect,destination):
    pdf=fitz.open(path);page=pdf[page_index]
    pix=page.get_pixmap(matrix=fitz.Matrix(1.6,1.6),clip=rect)
    image=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    image.thumbnail((920,440));out=Image.new('RGB',(960,540),'#f2f0eb');out.paste(image,((960-image.width)//2,50))
    draw=ImageDraw.Draw(out);draw.rectangle((0,0,960,42),fill='#536145');draw.text((20,9),'PRATINJAU / POTONGAN FILE ASLI / BUKAN FILE LENGKAP',font=font,fill='white')
    draw.text((20,505),'Ranel Barber v1.0 - contoh fiktif / lembar kerja',font=font,fill='#253027')
    out.save(destination,format='WEBP',quality=88)
    ledger.append({'path':str(destination.relative_to(ROOT)),'source':str(path.relative_to(ROOT)),'page':page_index+1,'clip':list(rect),'sha256':hashlib.sha256(destination.read_bytes()).hexdigest()})

for p in registry:
    folder=ROOT/p['path'];bundle=folder.name;artifact=release/(bundle+'.zip');expected=p['files']
    with zipfile.ZipFile(artifact,'w',zipfile.ZIP_DEFLATED) as z:
        for f in expected:z.write(folder/f,f'{bundle}/{f}')
    extracted=ROOT/'private-products/extracted'/bundle
    if extracted.exists():shutil.rmtree(extracted)
    extracted.mkdir(parents=True)
    with zipfile.ZipFile(artifact) as z:
        assert z.testzip() is None
        assert z.namelist()==[f'{bundle}/{f}' for f in expected]
        assert all(not n.startswith('/') and '..' not in Path(n).parts for n in z.namelist())
        z.extractall(extracted)
    content=extracted/bundle
    assert set(x.name for x in content.iterdir())==set(expected)
    listed=re.findall(r'^- `([^`]+)`$',(content/'MANIFEST.md').read_text(),re.M);assert listed==expected
    for f in expected:
        path=content/f;assert path.read_bytes()==(folder/f).read_bytes()
        if path.suffix=='.pdf':assert len(fitz.open(path))>0
        elif path.suffix=='.docx':assert Document(path).tables
        elif path.suffix=='.xlsx':
            w=load_workbook(path,data_only=False);cached=load_workbook(path,data_only=True)
            assert any(c.data_type=='f' for ws in w for row in ws for c in row)
            assert not any(c.data_type=='e' for ws in cached for row in ws for c in row)
    p['artifactSha256']=hashlib.sha256(artifact.read_bytes()).hexdigest();p['artifactBytes']=artifact.stat().st_size
    p['assetState']='PRODUCT_READY_FOR_COMMERCE_INTEGRATION';p['privateSource']=True
    slug=p['id'].split('.')[-1]
    crop(folder/'00-README.pdf',0,fitz.Rect(35,45,560,285),preview/(slug+'-quickstart.webp'))
    out=ROOT/'qa-artifacts/product-maturation-qc'/bundle
    if slug=='growth':
        path=out/'rendered-xlsx/05-Rekap-Kunjungan.pdf';d=fitz.open(path)
        index=next(i for i,page in enumerate(d) if 'DATA CONTOH SINTETIS SAJA' in page.get_text())
        crop(path,index,fitz.Rect(20,20,d[index].rect.width-20,340),preview/'growth-spreadsheet.webp')
    else:
        doc='03-Checklist-Buka-Tutup.pdf' if slug=='starter' else '09-Peta-Operating-System.pdf'
        crop(out/'rendered-docx'/doc,1,fitz.Rect(35,45,560,285),preview/(slug+'-document.webp' if slug=='starter' else 'system-map.webp'))
    p['previews']=[x['path'] for x in ledger if x['path'].split('/')[-1].startswith(slug+'-')]
qc['state']='PRODUCT_READY_FOR_COMMERCE_INTEGRATION'
qc['checks'].extend(['All workbook print views render and text stays in page bounds','Clean ZIP extraction and exact entry allowlist','Extracted PDF/DOCX/XLSX open; bytes and manifest match QC sources','Real-file cropped preview provenance','Visual contact-sheet review of all 82 rendered pages; thumbnail fine-detail limits'])
qc['limitations'].extend(['Repository is PUBLIC: matured paid sources/ZIPs are Git-ignored; historical v1.0 sources remain in public Git history.','Content readiness is not merchant/checkout/secure delivery or legal/tax clearance.'])
qc['archives']=[{'path':p['releaseArtifact'],'sha256':p['artifactSha256'],'bytes':p['artifactBytes'],'files':p['files']} for p in registry]
qc['previews']=ledger
(ROOT/'private-products/registry.json').write_text(json.dumps(registry,ensure_ascii=False,indent=2)+'\n')
(ROOT/'products/registry.json').write_text(json.dumps(registry,ensure_ascii=False,indent=2)+'\n')
(ROOT/'docs/products/product-maturation-qc.json').write_text(json.dumps(qc,ensure_ascii=False,indent=2)+'\n')
print('ZIP/extraction/manifest/hash/open checks PASS: 3 archives, 31 files. Six cropped actual previews generated.')
